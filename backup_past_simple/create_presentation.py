import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(is_dark=False, filename="Past_Simple_Tense_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    if is_dark:
        # Dark Theme Palette
        BG_SLIDE = RGBColor(15, 23, 42)        # #0f172a Deep Slate
        BG_CARD = RGBColor(30, 41, 59)         # #1e293b Slate 800
        BORDER_COLOR = RGBColor(51, 65, 85)    # #334155 Slate 700
        PRIMARY = RGBColor(96, 165, 250)       # #60a5fa Bright Blue
        SECONDARY = RGBColor(56, 189, 248)     # #38bdf8 Sky 400
        ACCENT = RGBColor(251, 191, 36)        # #fbbf24 Amber
        SUCCESS = RGBColor(74, 222, 128)       # #4ade80 Emerald
        SUCCESS_DARK = RGBColor(74, 222, 128)
        DANGER = RGBColor(248, 113, 113)       # #f87171 Red 400
        TEXT_DARK = RGBColor(241, 245, 249)    # #f1f5f9
        TEXT_MUTED = RGBColor(148, 163, 184)   # #94a3b8
        WHITE = RGBColor(255, 255, 255)
        LIGHT_BLUE = RGBColor(23, 37, 84)      # Dark Blue Box
        LIGHT_GREEN = RGBColor(6, 78, 59)      # Dark Green Box
        LIGHT_RED = RGBColor(69, 10, 10)       # Dark Red Box
        LIGHT_AMBER = RGBColor(69, 26, 3)      # Dark Amber Box
        TBL_HEADER_BG = RGBColor(30, 58, 138)
        TBL_ROW_ALT = RGBColor(15, 23, 42)
    else:
        # Light Theme Palette
        BG_SLIDE = RGBColor(255, 255, 255)
        BG_CARD = RGBColor(248, 250, 252)
        BORDER_COLOR = RGBColor(226, 232, 240)
        PRIMARY = RGBColor(30, 58, 138)
        SECONDARY = RGBColor(2, 132, 199)
        ACCENT = RGBColor(217, 119, 6)
        SUCCESS = RGBColor(22, 163, 74)
        SUCCESS_DARK = RGBColor(21, 128, 61)
        DANGER = RGBColor(220, 38, 38)
        TEXT_DARK = RGBColor(30, 41, 59)
        TEXT_MUTED = RGBColor(100, 116, 139)
        WHITE = RGBColor(255, 255, 255)
        LIGHT_BLUE = RGBColor(239, 246, 255)
        LIGHT_GREEN = RGBColor(240, 253, 244)
        LIGHT_RED = RGBColor(254, 242, 242)
        LIGHT_AMBER = RGBColor(254, 243, 199)
        TBL_HEADER_BG = RGBColor(241, 245, 249)
        TBL_ROW_ALT = BG_CARD

    FONT_MAIN = "Sarabun"
    FONT_HEADING = "Outfit"

    def apply_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_SLIDE
        bg.line.fill.background()

    def add_header(slide, tag_text, title_text, slide_num):
        apply_slide_bg(slide)

        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = tag_text.upper()
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = SECONDARY

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY if not is_dark else SECONDARY
        p1.space_before = Pt(2)

        num_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.45), Inches(2.0), Inches(0.5))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = f"Slide {slide_num} / 10"
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(13)
        p_num.font.color.rgb = TEXT_MUTED

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_COLOR
        line.line.color.rgb = BORDER_COLOR

    # -------------------------------------------------------------
    # SLIDE 1: Cover
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(15, 23, 42) if is_dark else RGBColor(30, 58, 138)
    bg1.line.fill.background()

    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.666), Inches(1.6), Inches(4.0), Inches(0.5))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = RGBColor(37, 99, 235) if is_dark else SECONDARY
    badge1.line.fill.background()
    tf_b1 = badge1.text_frame
    p_b1 = tf_b1.paragraphs[0]
    p_b1.text = "ENGLISH GRAMMAR • GRADE 8 (ม.2)"
    p_b1.alignment = PP_ALIGN.CENTER
    p_b1.font.name = FONT_HEADING
    p_b1.font.size = Pt(12)
    p_b1.font.bold = True
    p_b1.font.color.rgb = WHITE

    tbox1 = s1.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(10.333), Inches(3.2))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    p_t1 = tf1.paragraphs[0]
    p_t1.text = "Past Simple Tense"
    p_t1.alignment = PP_ALIGN.CENTER
    p_t1.font.name = FONT_HEADING
    p_t1.font.size = Pt(54)
    p_t1.font.bold = True
    p_t1.font.color.rgb = WHITE if not is_dark else RGBColor(56, 189, 248)

    p_sub1 = tf1.add_paragraph()
    p_sub1.text = "สรุปหลักไวยากรณ์ กฎการสร้างประโยค และเทคนิคจำแม่นเพื่อการสื่อสารและทำข้อสอบ"
    p_sub1.alignment = PP_ALIGN.CENTER
    p_sub1.font.name = FONT_MAIN
    p_sub1.font.size = Pt(20)
    p_sub1.font.color.rgb = RGBColor(224, 242, 254)
    p_sub1.space_before = Pt(18)

    p_note1 = tf1.add_paragraph()
    p_note1.text = "ฉบับคู่มือการสอนติวเข้มสำหรับลูก (Unit 5 & Unit 6)"
    p_note1.alignment = PP_ALIGN.CENTER
    p_note1.font.name = FONT_MAIN
    p_note1.font.size = Pt(15)
    p_note1.font.color.rgb = RGBColor(186, 230, 253)
    p_note1.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 2: Core Concept
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, "Core Concept", "หัวใจหลักของ Past Simple Tense", 2)

    b_concept = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.4))
    b_concept.fill.solid()
    b_concept.fill.fore_color.rgb = LIGHT_BLUE
    b_concept.line.color.rgb = RGBColor(30, 58, 138) if is_dark else RGBColor(191, 219, 254)
    tf_c = b_concept.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = Inches(0.3)
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "💡 ใช้เมื่อไหร่?"
    p_c1.font.name = FONT_MAIN
    p_c1.font.size = Pt(16)
    p_c1.font.bold = True
    p_c1.font.color.rgb = PRIMARY
    p_c2 = tf_c.add_paragraph()
    p_c2.text = 'ใช้กับเหตุการณ์หรือการกระทำที่ "เกิดขึ้นและจบสิ้นลงอย่างสมบูรณ์แล้วในอดีต" (ปัจจุบันไม่ได้ทำแล้ว)'
    p_c2.font.name = FONT_MAIN
    p_c2.font.size = Pt(15)
    p_c2.font.color.rgb = TEXT_DARK
    p_c2.space_before = Pt(6)

    tb_time_hdr = s2.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(11.733), Inches(0.5))
    p_th = tb_time_hdr.text_frame.paragraphs[0]
    p_th.text = "⏰ คำบอกเวลาที่พบบ่อย (Past Time Expressions)"
    p_th.font.name = FONT_MAIN
    p_th.font.size = Pt(16)
    p_th.font.bold = True
    p_th.font.color.rgb = PRIMARY

    time_cards = [
        ("yesterday", "เมื่อวานนี้\n(yesterday morning, yesterday afternoon)"),
        ("last + เวลา", "...ที่แล้ว\n(last night, last week, last month, last year)"),
        ("ช่วงเวลา + ago", "...ที่ผ่านมา\n(two hours ago, three days ago, a year ago)"),
        ("in + ปีอดีต", "ในปี...\n(in 2020, in the past, when I was young)")
    ]

    left_pos = Inches(0.8)
    card_width = Inches(2.78)
    for title, desc in time_cards:
        c_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(4.1), card_width, Inches(2.7))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = BG_CARD
        c_shape.line.color.rgb = BORDER_COLOR
        tf_card = c_shape.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = tf_card.margin_right = Inches(0.2)
        p_ct = tf_card.paragraphs[0]
        p_ct.text = title
        p_ct.font.name = FONT_HEADING
        p_ct.font.size = Pt(16)
        p_ct.font.bold = True
        p_ct.font.color.rgb = SECONDARY
        p_cd = tf_card.add_paragraph()
        p_cd.text = desc
        p_cd.font.name = FONT_MAIN
        p_cd.font.size = Pt(13)
        p_cd.font.color.rgb = TEXT_DARK
        p_cd.space_before = Pt(8)
        left_pos += card_width + Inches(0.2)

    # -------------------------------------------------------------
    # SLIDE 3: Sentence Structures
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, "Grammar Rules", "โครงสร้างประโยค 3 รูปแบบ (+, -, ?)", 3)

    col_w = Inches(3.7)
    cols = [
        ("1. บอกเล่า (+)", "Subject + V.2", "ใช้กริยาช่องที่ 2 เสมอ", "• We played football yesterday.\n• He went to Japan last year.", SUCCESS, LIGHT_GREEN),
        ("2. ปฏิเสธ (-)", "S + didn't + V.inf", "ใช้ didn't แล้วคืนรูปกริยาเดิม", "• We didn't play football.\n• He didn't go to Japan.", DANGER, LIGHT_RED),
        ("3. คำถาม (?)", "Did + S + V.inf ...?", "เอา Did ขึ้นหน้า กริยาใช้รูปเดิม", "• Did they play football?\n• Did he go to Japan?", SECONDARY, LIGHT_BLUE)
    ]

    c_left = Inches(0.8)
    for title, formula, note, examples, theme_color, bg_box in cols:
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(1.8), col_w, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR
        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.25)

        p1 = tf_c.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = theme_color

        p2 = tf_c.add_paragraph()
        p2.text = f"  {formula}  "
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = theme_color
        p2.space_before = Pt(12)

        p3 = tf_c.add_paragraph()
        p3.text = note
        p3.font.name = FONT_MAIN
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(10)

        p4 = tf_c.add_paragraph()
        p4.text = "ตัวอย่าง:"
        p4.font.name = FONT_MAIN
        p4.font.size = Pt(13)
        p4.font.bold = True
        p4.font.color.rgb = TEXT_DARK
        p4.space_before = Pt(16)

        p5 = tf_c.add_paragraph()
        p5.text = examples
        p5.font.name = FONT_MAIN
        p5.font.size = Pt(13)
        p5.font.color.rgb = TEXT_DARK
        p5.space_before = Pt(6)

        c_left += col_w + Inches(0.316)

    # -------------------------------------------------------------
    # SLIDE 4: Regular Verbs Rules
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header(s4, "Regular Verbs", "กฎ 5 ข้อของการเติม -ed (Regular Verbs)", 4)

    rows, cols_cnt = 6, 3
    tbl_shape = s4.shapes.add_table(rows, cols_cnt, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(5.0)
    tbl.columns[2].width = Inches(3.533)

    headers = ["กฎการเติม", "คำอธิบาย", "ตัวอย่าง"]
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TBL_HEADER_BG
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_MAIN
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE if is_dark else PRIMARY

    table_data = [
        ("1. เติม -ed ทั่วไป", "กริยาส่วนใหญ่เติม -ed ได้ทันที", "watch → watched\nclean → cleaned"),
        ("2. ลงท้ายด้วย e", "เติมเพียง -d ท้ายคำได้เลย", "live → lived\narrive → arrived"),
        ("3. พยัญชนะ + y", "เปลี่ยน y เป็น i แล้วเติม -ed", "study → studied\ntidy → tidied"),
        ("4. สระ + y", "เติม -ed ได้ทันที (ห้ามเปลี่ยนรูป)", "play → played\nstay → stayed"),
        ("5. 1 สระ + 1 ตัวสะกด", "เบิ้ลตัวสะกดท้ายอีก 1 ตัวก่อนเติม -ed", "stop → stopped\nplan → planned")
    ]

    for row_idx, data in enumerate(table_data, start=1):
        for col_idx, text in enumerate(data):
            cell = tbl.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_CARD if row_idx % 2 != 0 else TBL_ROW_ALT
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.name = FONT_MAIN
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_DARK
            if col_idx == 0:
                p.font.bold = True

    # -------------------------------------------------------------
    # SLIDE 5: Irregular Verbs Table
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header(s5, "Irregular Verbs", "กริยาเปลี่ยนรูปยอดฮิต (ต้องจำให้แม่น)", 5)

    t1_data = [
        ("Base (V.1)", "Past (V.2)", "ความหมาย"),
        ("eat", "ate", "กิน"),
        ("buy", "bought", "ซื้อ"),
        ("go", "went", "ไป"),
        ("see", "saw", "เห็น / พบ"),
        ("have", "had", "มี / รับประทาน")
    ]

    t2_data = [
        ("Base (V.1)", "Past (V.2)", "ความหมาย"),
        ("make", "made", "ทำ / สร้าง"),
        ("lose", "lost", "ทำหาย / แพ้"),
        ("read", "read (เรด)", "อ่าน"),
        ("write", "wrote", "เขียน"),
        ("drink", "drank", "ดื่ม")
    ]

    def render_verb_tbl(slide, data, left, top, width, height):
        tbl_s = slide.shapes.add_table(len(data), 3, left, top, width, height)
        t = tbl_s.table
        t.columns[0].width = Inches(1.8)
        t.columns[1].width = Inches(2.0)
        t.columns[2].width = width - Inches(3.8)
        for r_idx, row in enumerate(data):
            for c_idx, val in enumerate(row):
                cell = t.cell(r_idx, c_idx)
                cell.fill.solid()
                if r_idx == 0:
                    cell.fill.fore_color.rgb = TBL_HEADER_BG
                else:
                    cell.fill.fore_color.rgb = BG_CARD if r_idx % 2 != 0 else TBL_ROW_ALT
                p = cell.text_frame.paragraphs[0]
                p.text = val
                p.font.name = FONT_MAIN
                p.font.size = Pt(13 if r_idx > 0 else 14)
                p.font.bold = (r_idx == 0 or c_idx < 2)
                p.font.color.rgb = (WHITE if is_dark else PRIMARY) if r_idx == 0 else (SUCCESS_DARK if c_idx == 1 else TEXT_DARK)

    render_verb_tbl(s5, t1_data, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8))
    render_verb_tbl(s5, t2_data, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))

    # -------------------------------------------------------------
    # SLIDE 6: Critical Warning
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_header(s6, "Common Trap", "⚠️ กฎทอง \"ถอดเครื่องแบบ\" (จุดตายของข้อสอบ)", 6)

    w_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.6))
    w_box.fill.solid()
    w_box.fill.fore_color.rgb = LIGHT_AMBER
    w_box.line.color.rgb = ACCENT
    tf_w = w_box.text_frame
    tf_w.word_wrap = True
    tf_w.margin_left = tf_w.margin_right = Inches(0.3)
    p_w1 = tf_w.paragraphs[0]
    p_w1.text = "เมื่อไหร่ที่มี Did หรือ didn't โผล่เข้ามาในประโยค..."
    p_w1.font.name = FONT_MAIN
    p_w1.font.size = Pt(16)
    p_w1.font.bold = True
    p_w1.font.color.rgb = RGBColor(252, 211, 77) if is_dark else RGBColor(146, 64, 14)
    p_w2 = tf_w.add_paragraph()
    p_w2.text = "กริยาแท้ต้องถูก \"ถอดรูปอดีตทิ้ง\" แล้วกลับไปเป็น กริยาช่อง 1 ตัวเดิม (Base Form / V.inf) เสมอ!"
    p_w2.font.name = FONT_MAIN
    p_w2.font.size = Pt(15)
    p_w2.font.bold = True
    p_w2.font.color.rgb = RGBColor(252, 211, 77) if is_dark else RGBColor(146, 64, 14)
    p_w2.space_before = Pt(6)

    col_width_half = Inches(5.7)
    b_wrong = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.7), col_width_half, Inches(3.2))
    b_wrong.fill.solid()
    b_wrong.fill.fore_color.rgb = LIGHT_RED
    b_wrong.line.color.rgb = DANGER
    tf_wr = b_wrong.text_frame
    tf_wr.word_wrap = True
    tf_wr.margin_left = tf_wr.margin_right = Inches(0.3)
    p_wr1 = tf_wr.paragraphs[0]
    p_wr1.text = "❌ สิ่งที่เด็กๆ ชอบผิด"
    p_wr1.font.name = FONT_MAIN
    p_wr1.font.size = Pt(17)
    p_wr1.font.bold = True
    p_wr1.font.color.rgb = DANGER
    p_wr2 = tf_wr.add_paragraph()
    p_wr2.text = "• He didn't went to school. ❌\n• Did you saw that cat? ❌\n• We didn't watched TV. ❌"
    p_wr2.font.name = FONT_MAIN
    p_wr2.font.size = Pt(15)
    p_wr2.font.color.rgb = TEXT_DARK
    p_wr2.space_before = Pt(14)

    b_corr = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(3.7), col_width_half, Inches(3.2))
    b_corr.fill.solid()
    b_corr.fill.fore_color.rgb = LIGHT_GREEN
    b_corr.line.color.rgb = SUCCESS
    tf_cr = b_corr.text_frame
    tf_cr.word_wrap = True
    tf_cr.margin_left = tf_cr.margin_right = Inches(0.3)
    p_cr1 = tf_cr.paragraphs[0]
    p_cr1.text = "✔️ ประโยคที่ถูกต้อง"
    p_cr1.font.name = FONT_MAIN
    p_cr1.font.size = Pt(17)
    p_cr1.font.bold = True
    p_cr1.font.color.rgb = SUCCESS
    p_cr2 = tf_cr.add_paragraph()
    p_cr2.text = "• He didn't go to school. ✔️\n• Did you see that cat? ✔️\n• We didn't watch TV. ✔️"
    p_cr2.font.name = FONT_MAIN
    p_cr2.font.size = Pt(15)
    p_cr2.font.color.rgb = TEXT_DARK
    p_cr2.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 7: Step-by-Step Sentence Building
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_header(s7, "Skill Building", "สูตรกันตาย 3 สเต็ปสร้างประโยคจาก Verb", 7)

    tb_intro = s7.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    p_in = tb_intro.text_frame.paragraphs[0]
    p_in.text = "เมื่อโจทย์ให้ Verb มาคำเดียว แล้วคิดไม่ออก ให้จำ 3 สเต็ปนี้:"
    p_in.font.name = FONT_MAIN
    p_in.font.size = Pt(15)
    p_in.font.color.rgb = TEXT_DARK

    steps = [
        ("Step 1", "ผันเป็น V.2 ให้ถูก", "เปลี่ยนรูป หรือเติม -ed ทันที", "eat → ate", SECONDARY),
        ("Step 2", "เติมนามสั้นๆ 1 ตัว", "ทำอะไร? หรือไปที่ไหน?", "กินอะไร? → noodles", SECONDARY),
        ("Step 3", "ประกอบร่างเข้าสูตร", "I + [V.2] + [กรรม] + yesterday.", "I ate noodles yesterday.", SUCCESS)
    ]

    s_left = Inches(0.8)
    s_width = Inches(3.7)
    for step_num, step_title, step_desc, step_ex, st_color in steps:
        s_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, s_left, Inches(2.2), s_width, Inches(3.8))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = BG_CARD
        s_box.line.color.rgb = BORDER_COLOR
        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_right = Inches(0.25)
        tf_s.margin_top = Inches(0.25)

        p1 = tf_s.paragraphs[0]
        p1.text = step_num
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = st_color

        p2 = tf_s.add_paragraph()
        p2.text = step_title
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = WHITE if is_dark else PRIMARY
        p2.space_before = Pt(6)

        p3 = tf_s.add_paragraph()
        p3.text = step_desc
        p3.font.name = FONT_MAIN
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(8)

        p4 = tf_s.add_paragraph()
        p4.text = f"ตัวอย่าง:\n{step_ex}"
        p4.font.name = FONT_MAIN
        p4.font.size = Pt(14)
        p4.font.bold = True
        p4.font.color.rgb = st_color
        p4.space_before = Pt(16)

        s_left += s_width + Inches(0.316)

    wb_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.85))
    wb_box.fill.solid()
    wb_box.fill.fore_color.rgb = LIGHT_GREEN
    wb_box.line.color.rgb = SUCCESS
    tf_wb = wb_box.text_frame
    p_wb = tf_wb.paragraphs[0]
    p_wb.text = "💡 คลังคำนามช่วยชีวิต: noodles, milk, pizza, my room, YouTube, football, to school, a cat, a movie"
    p_wb.font.name = FONT_MAIN
    p_wb.font.size = Pt(13)
    p_wb.font.bold = True
    p_wb.font.color.rgb = SUCCESS_DARK

    # -------------------------------------------------------------
    # SLIDE 8: -ed vs -ing Adjectives
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_header(s8, "Unit 5 & 6 Focus", "Adjectives: ลงท้าย -ed VS -ing", 8)

    adj_w = Inches(5.7)
    ed_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), adj_w, Inches(5.1))
    ed_card.fill.solid()
    ed_card.fill.fore_color.rgb = LIGHT_GREEN
    ed_card.line.color.rgb = SUCCESS
    tf_ed = ed_card.text_frame
    tf_ed.word_wrap = True
    tf_ed.margin_left = tf_ed.margin_right = Inches(0.3)
    tf_ed.margin_top = Inches(0.3)

    p_e1 = tf_ed.paragraphs[0]
    p_e1.text = "😊 ลงท้ายด้วย -ed = \"รู้สึก...\""
    p_e1.font.name = FONT_MAIN
    p_e1.font.size = Pt(18)
    p_e1.font.bold = True
    p_e1.font.color.rgb = SUCCESS

    p_e2 = tf_ed.add_paragraph()
    p_e2.text = "ใช้อธิบาย ความรู้สึกของคนหรือสัตว์"
    p_e2.font.name = FONT_MAIN
    p_e2.font.size = Pt(14)
    p_e2.font.color.rgb = TEXT_DARK
    p_e2.space_before = Pt(8)

    p_e3 = tf_ed.add_paragraph()
    p_e3.text = "• bored = รู้สึกเบื่อ\n• interested = รู้สึกสนใจ\n• relaxed = รู้สึกผ่อนคลาย\n• annoyed = รู้สึกหงุดหงิด/รำคาญ"
    p_e3.font.name = FONT_MAIN
    p_e3.font.size = Pt(14)
    p_e3.font.color.rgb = TEXT_DARK
    p_e3.space_before = Pt(12)

    p_e4 = tf_ed.add_paragraph()
    p_e4.text = 'ตัวอย่าง:\n"I was bored during the lesson."'
    p_e4.font.name = FONT_MAIN
    p_e4.font.size = Pt(14)
    p_e4.font.bold = True
    p_e4.font.color.rgb = SUCCESS
    p_e4.space_before = Pt(16)

    ing_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.8), adj_w, Inches(5.1))
    ing_card.fill.solid()
    ing_card.fill.fore_color.rgb = LIGHT_RED
    ing_card.line.color.rgb = DANGER
    tf_ing = ing_card.text_frame
    tf_ing.word_wrap = True
    tf_ing.margin_left = tf_ing.margin_right = Inches(0.3)
    tf_ing.margin_top = Inches(0.3)

    p_i1 = tf_ing.paragraphs[0]
    p_i1.text = "🎬 ลงท้ายด้วย -ing = \"น่า...\""
    p_i1.font.name = FONT_MAIN
    p_i1.font.size = Pt(18)
    p_i1.font.bold = True
    p_i1.font.color.rgb = DANGER

    p_i2 = tf_ing.add_paragraph()
    p_i2.text = "ใช้อธิบาย ลักษณะของสิ่งของ หนัง หรือเหตุการณ์"
    p_i2.font.name = FONT_MAIN
    p_i2.font.size = Pt(14)
    p_i2.font.color.rgb = TEXT_DARK
    p_i2.space_before = Pt(8)

    p_i3 = tf_ing.add_paragraph()
    p_i3.text = "• boring = น่าเบื่อ\n• interesting = น่าสนใจ\n• relaxing = น่าผ่อนคลาย\n• annoying = น่ารำคาญ"
    p_i3.font.name = FONT_MAIN
    p_i3.font.size = Pt(14)
    p_i3.font.color.rgb = TEXT_DARK
    p_i3.space_before = Pt(12)

    p_i4 = tf_ing.add_paragraph()
    p_i4.text = 'ตัวอย่าง:\n"That movie was really boring."'
    p_i4.font.name = FONT_MAIN
    p_i4.font.size = Pt(14)
    p_i4.font.bold = True
    p_i4.font.color.rgb = DANGER
    p_i4.space_before = Pt(16)

    # -------------------------------------------------------------
    # SLIDE 9: Exam Traps
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_slide_layout)
    add_header(s9, "Exam Diagnostics", "3 จุดหลอกยอดฮิตที่เด็ก ม.2 ชอบเสียคะแนน", 9)

    traps = [
        ("1. สะกดคำผิดเมื่อเติม -ed กับตัว y", "• ถ้าหน้า y เป็นพยัญชนะ เปลี่ยน y → i + ed เช่น studied (ไม่ใช่ studyed ❌)\n• ถ้าหน้า y เป็นสระ เติม -ed ได้เลย เช่น played (ไม่ใช่ plaied ❌)"),
        ("2. คำบอกเวลากลางคืน", "• ภาษาอังกฤษใช้ last night เท่านั้น\n• ห้ามใช้คำว่า yesterday night ❌ โดยเด็ดขาด"),
        ("3. การแสดงความเป็นเจ้าของ Double Genitive", "• โครงสร้าง a friend of + mine / yours / his / hers / name's\n• ตัวอย่าง: a friend of mine (ไม่ใช่ a friend of me ❌)")
    ]

    t_top = Inches(1.8)
    for title, desc in traps:
        t_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), t_top, Inches(11.733), Inches(1.5))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = BG_CARD
        t_card.line.color.rgb = BORDER_COLOR
        tf_t = t_card.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = Inches(0.3)
        tf_t.margin_top = Inches(0.2)

        p1 = tf_t.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = SECONDARY if is_dark else PRIMARY

        p2 = tf_t.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(6)

        t_top += Inches(1.7)

    # -------------------------------------------------------------
    # SLIDE 10: Closing Slide
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_slide_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = RGBColor(19, 78, 74) if is_dark else RGBColor(15, 118, 110)
    bg10.line.fill.background()

    badge10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.166), Inches(1.6), Inches(3.0), Inches(0.5))
    badge10.fill.solid()
    badge10.fill.fore_color.rgb = RGBColor(13, 148, 136)
    badge10.line.fill.background()
    tf_b10 = badge10.text_frame
    p_b10 = tf_b10.paragraphs[0]
    p_b10.text = "GREAT JOB!"
    p_b10.alignment = PP_ALIGN.CENTER
    p_b10.font.name = FONT_HEADING
    p_b10.font.size = Pt(13)
    p_b10.font.bold = True
    p_b10.font.color.rgb = WHITE

    tbox10 = s10.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(10.333), Inches(3.5))
    tf10 = tbox10.text_frame
    tf10.word_wrap = True
    p_t10 = tf10.paragraphs[0]
    p_t10.text = "พร้อมลุยทำแบบฝึกหัดแล้ว!"
    p_t10.alignment = PP_ALIGN.CENTER
    p_t10.font.name = FONT_MAIN
    p_t10.font.size = Pt(46)
    p_t10.font.bold = True
    p_t10.font.color.rgb = WHITE

    p_sub10 = tf10.add_paragraph()
    p_sub10.text = '"ฝึกฝนสม่ำเสมอ จับตาดู Did/didn\'t และจำรูปกริยาให้แม่น แล้วลูกจะทำข้อสอบได้อย่างมั่นใจแน่นอน"'
    p_sub10.alignment = PP_ALIGN.CENTER
    p_sub10.font.name = FONT_MAIN
    p_sub10.font.size = Pt(18)
    p_sub10.font.color.rgb = RGBColor(204, 251, 241)
    p_sub10.space_before = Pt(18)

    p_note10 = tf10.add_paragraph()
    p_note10.text = "คุณพ่อสามารถเปิดหน้าแบบฝึกหัดให้น้องเริ่มทำต่อได้เลยครับ 👏"
    p_note10.alignment = PP_ALIGN.CENTER
    p_note10.font.name = FONT_MAIN
    p_note10.font.size = Pt(15)
    p_note10.font.color.rgb = RGBColor(153, 246, 228)
    p_note10.space_before = Pt(14)

    prs.save(filename)
    print(f"Saved {filename}")

if __name__ == "__main__":
    create_deck(is_dark=False, filename="Past_Simple_Tense_Presentation.pptx")
    create_deck(is_dark=True, filename="Past_Simple_Tense_Presentation_Dark.pptx")
